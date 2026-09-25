#!/usr/bin/env python3
"""
MT5 交易整理 → Excel（依作業 / EA 分組計算損益）

在「正在跑 MT5 的那台 Windows 電腦」上執行：
    pip install MetaTrader5 openpyxl
    python scripts/mt5_trades_report.py --from 2026-09-22 --strategies v5-3 STT5M

它會：
  1. 連到已開啟並登入的 MT5 終端，抓 --from 起到現在的所有成交 (deals)
  2. 以 position_id 把進場 / 出場成交合併成「一張單」
  3. 依開倉單的 comment（或 magic number）判斷屬於哪個作業
  4. 輸出 Excel：總覽（每個作業的筆數、勝率、總虧損、淨損益、最大回撤…）、
     每個作業一個分頁、未平倉分頁

作業判斷規則（依序）：
  --map 12345=v5-3        magic number 直接對應作業名稱（最準）
  --strategies v5-3 STT5M 開倉 comment 含有此字串（不分大小寫）就歸到該作業
  否則                    用開倉 comment 本身；comment 空白則為 magic_<magic>

也可以不連 MT5，改讀本腳本 --dump 出來的 CSV：
    python scripts/mt5_trades_report.py --csv deals.csv --strategies v5-3 STT5M
"""
import argparse
import csv
import re
import sys
from collections import OrderedDict, defaultdict
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# MT5 enums (MetaTrader5.DEAL_TYPE_* / DEAL_ENTRY_*)
DEAL_TYPE_BUY, DEAL_TYPE_SELL = 0, 1
ENTRY_IN, ENTRY_OUT, ENTRY_INOUT, ENTRY_OUT_BY = 0, 1, 2, 3

DEAL_FIELDS = ["ticket", "order", "position_id", "time", "type", "entry", "magic",
               "symbol", "volume", "price", "profit", "commission", "swap", "fee", "comment"]


# ---------------------------------------------------------------- data sources

def load_from_mt5(start, end, login=None, password=None, server=None, path=None):
    try:
        import MetaTrader5 as mt5
    except ImportError:
        sys.exit("找不到 MetaTrader5 套件，請先執行: pip install MetaTrader5（僅支援 Windows）")

    kwargs = {}
    if path:
        kwargs["path"] = path
    if login:
        kwargs.update(login=int(login), password=password or "", server=server or "")
    if not mt5.initialize(**kwargs):
        sys.exit(f"MT5 連線失敗: {mt5.last_error()}（請確認 MT5 已開啟並登入，且允許演算法交易）")

    try:
        acc = mt5.account_info()
        raw = mt5.history_deals_get(start, end)
        if raw is None:
            sys.exit(f"讀取歷史成交失敗: {mt5.last_error()}")
        deals = [{f: getattr(d, f) for f in DEAL_FIELDS} for d in raw]

        # 往前多抓一點，找回「9/22 前開倉、之後才平倉」單的進場成交，避免歸類錯誤
        need = {d["position_id"] for d in deals
                if d["entry"] in (ENTRY_OUT, ENTRY_OUT_BY)
                and not any(x["position_id"] == d["position_id"] and x["entry"] == ENTRY_IN for x in deals)}
        for pid in need:
            for d in mt5.history_deals_get(position=pid) or []:
                if d.ticket not in {x["ticket"] for x in deals}:
                    deals.append({f: getattr(d, f) for f in DEAL_FIELDS})

        open_pos = [{
            "position_id": p.identifier, "symbol": p.symbol, "type": p.type, "volume": p.volume,
            "price_open": p.price_open, "price_current": p.price_current, "time": p.time,
            "sl": p.sl, "tp": p.tp, "profit": p.profit, "swap": p.swap,
            "magic": p.magic, "comment": p.comment,
        } for p in (mt5.positions_get() or [])]
        account = f"{acc.login} @ {acc.server} ({acc.currency})" if acc else ""
        currency = acc.currency if acc else ""
    finally:
        mt5.shutdown()
    return deals, open_pos, account, currency


def load_from_csv(path):
    deals = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            d = dict(row)
            for k in ("ticket", "order", "position_id", "time", "type", "entry", "magic"):
                d[k] = int(float(d[k] or 0))
            for k in ("volume", "price", "profit", "commission", "swap", "fee"):
                d[k] = float(d[k] or 0)
            deals.append(d)
    return deals


def dump_csv(deals, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=DEAL_FIELDS)
        w.writeheader()
        for d in sorted(deals, key=lambda x: (x["time"], x["ticket"])):
            w.writerow({k: d.get(k, "") for k in DEAL_FIELDS})


# ---------------------------------------------------------------- processing

def classify(magic, comment, magic_map, patterns):
    if magic in magic_map:
        return magic_map[magic]
    c = (comment or "").strip()
    for name in patterns:
        if name.lower() in c.lower():
            return name
    if c and not re.match(r"^(sl|tp|so)\b|^\[", c, re.I):
        return c
    return f"magic_{magic}"


def build_positions(deals, start_ts, magic_map, patterns):
    """Group deals by position_id → one row per closed position opened at/after start."""
    by_pos = defaultdict(list)
    for d in deals:
        if d["type"] in (DEAL_TYPE_BUY, DEAL_TYPE_SELL) and d["position_id"]:
            by_pos[d["position_id"]].append(d)

    rows = []
    for pid, ds in by_pos.items():
        ds.sort(key=lambda x: (x["time"], x["ticket"]))
        ins = [d for d in ds if d["entry"] in (ENTRY_IN, ENTRY_INOUT)]
        outs = [d for d in ds if d["entry"] in (ENTRY_OUT, ENTRY_OUT_BY, ENTRY_INOUT)]
        if not ins or not outs:
            continue  # 仍未平倉，或缺進場資料
        first = ins[0]
        if first["time"] < start_ts:
            continue
        vol_in = sum(d["volume"] for d in ins)
        vol_out = sum(d["volume"] for d in outs)
        if vol_out + 1e-9 < vol_in:
            continue  # 部分平倉，仍有剩餘部位 → 算在未平倉
        wavg = lambda xs: sum(d["price"] * d["volume"] for d in xs) / max(sum(d["volume"] for d in xs), 1e-12)
        rows.append({
            "position_id": pid,
            "strategy": classify(first["magic"], first["comment"], magic_map, patterns),
            "symbol": first["symbol"],
            "side": "Buy" if first["type"] == DEAL_TYPE_BUY else "Sell",
            "volume": vol_in,
            "open_time": first["time"],
            "open_price": wavg(ins),
            "close_time": outs[-1]["time"],
            "close_price": wavg(outs),
            "profit": sum(d["profit"] for d in ds),
            "commission": sum(d["commission"] + d.get("fee", 0) for d in ds),
            "swap": sum(d["swap"] for d in ds),
            "comment": first["comment"],
            "close_comment": outs[-1]["comment"],
            "magic": first["magic"],
        })
    rows.sort(key=lambda r: (r["close_time"], r["position_id"]))
    return rows


# ---------------------------------------------------------------- excel

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
TOT_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
MONEY = '#,##0.00;[Red]-#,##0.00;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
TS = "yyyy-mm-dd hh:mm:ss"


def ts(t):
    # MT5 時間戳為「券商伺服器時間」，直接轉成不帶時區的 datetime 顯示
    return datetime.fromtimestamp(int(t), tz=timezone.utc).replace(tzinfo=None)


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def safe_sheet_name(name, used):
    s = re.sub(r"[\[\]:*?/\\]", "_", str(name))[:31] or "未命名"
    base, n = s, 2
    while s.lower() in used:
        s = f"{base[:28]}_{n}"
        n += 1
    used.add(s.lower())
    return s


STRAT_HEADERS = ["#", "持倉ID", "商品", "方向", "手數", "開倉時間", "開倉價", "平倉時間", "平倉價",
                 "毛損益", "手續費", "庫存費", "淨損益", "累計損益", "累計高點", "回撤",
                 "開倉註解", "平倉註解", "Magic"]


def write_strategy_sheet(ws, name, rows):
    ws.append(STRAT_HEADERS)
    style_header(ws, 1, len(STRAT_HEADERS))
    for i, r in enumerate(rows, 1):
        n = i + 1
        ws.append([
            i, r["position_id"], r["symbol"], r["side"], r["volume"],
            ts(r["open_time"]), r["open_price"], ts(r["close_time"]), r["close_price"],
            round(r["profit"], 2), round(r["commission"], 2), round(r["swap"], 2),
            f"=J{n}+K{n}+L{n}",
            f"=M{n}" if i == 1 else f"=N{n-1}+M{n}",
            f"=MAX(0,N{n})" if i == 1 else f"=MAX(O{n-1},N{n})",
            f"=N{n}-O{n}",
            r["comment"], r["close_comment"], r["magic"],
        ])
    last = len(rows) + 1
    tot = last + 1
    ws.cell(row=tot, column=1, value="合計")
    for col in ("J", "K", "L", "M"):
        ws[f"{col}{tot}"] = f"=SUM({col}2:{col}{last})"
    for c in range(1, len(STRAT_HEADERS) + 1):
        cell = ws.cell(row=tot, column=c)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = TOT_FILL
        cell.border = BORDER

    for row in ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.font = Font(name=FONT)
            cell.border = BORDER
    for row in ws.iter_rows(min_row=2, max_row=tot):
        for c in row:
            L = c.column_letter
            if L in "JKLMNOP":
                c.number_format = MONEY
            elif L in "FH":
                c.number_format = TS
            elif L in "GI":
                c.number_format = "0.00000"
            elif L == "E":
                c.number_format = "0.00"
    ws["F1"].comment = Comment("MT5 券商伺服器時間", "report")
    ws["K1"].comment = Comment("含 commission 與 fee", "report")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(STRAT_HEADERS))}{last}"
    set_widths(ws, [5, 12, 10, 6, 7, 19, 11, 19, 11, 11, 10, 10, 11, 11, 11, 11, 16, 16, 12])
    return last


SUMMARY_HEADERS = ["作業", "交易筆數", "獲利筆數", "虧損筆數", "勝率", "總獲利", "總虧損",
                   "淨損益", "平均獲利", "平均虧損", "最大單筆虧損", "獲利因子", "最大回撤",
                   "手續費", "庫存費"]


def write_summary(ws, groups, account, start, end, currency):
    ws["A1"] = "MT5 交易損益總覽（依作業分組）"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = f"帳戶: {account}" if account else "帳戶: (CSV 匯入)"
    ws["A3"] = f"期間: {start:%Y-%m-%d} ~ {end:%Y-%m-%d %H:%M}（以開倉時間計，只含已平倉單；金額單位 {currency or '帳戶幣別'}）"
    for a in ("A2", "A3"):
        ws[a].font = Font(name=FONT, color="595959")

    hr = 5
    for c, h in enumerate(SUMMARY_HEADERS, 1):
        ws.cell(row=hr, column=c, value=h)
    style_header(ws, hr, len(SUMMARY_HEADERS))

    r = hr
    for name, (sheet, last) in groups.items():
        r += 1
        q = f"'{sheet.replace(chr(39), chr(39) * 2)}'!"
        net = f"{q}$M$2:$M${last}"
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=1).hyperlink = f"#{q}A1"
        ws.cell(row=r, column=2, value=f"=COUNT({net})")
        ws.cell(row=r, column=3, value=f'=COUNTIF({net},">0")')
        ws.cell(row=r, column=4, value=f'=COUNTIF({net},"<0")')
        ws.cell(row=r, column=5, value=f"=IF(B{r}=0,0,C{r}/B{r})")
        ws.cell(row=r, column=6, value=f'=SUMIF({net},">0")')
        ws.cell(row=r, column=7, value=f'=SUMIF({net},"<0")')
        ws.cell(row=r, column=8, value=f"=SUM({net})")
        ws.cell(row=r, column=9, value=f"=IF(C{r}=0,0,F{r}/C{r})")
        ws.cell(row=r, column=10, value=f"=IF(D{r}=0,0,G{r}/D{r})")
        ws.cell(row=r, column=11, value=f"=MIN(0,MIN({net}))")
        ws.cell(row=r, column=12, value=f'=IF(G{r}=0,"",F{r}/-G{r})')
        ws.cell(row=r, column=13, value=f"=MIN(0,MIN({q}$P$2:$P${last}))")
        ws.cell(row=r, column=14, value=f"=SUM({q}$K$2:$K${last})")
        ws.cell(row=r, column=15, value=f"=SUM({q}$L$2:$L${last})")

    first, lastr = hr + 1, r
    tot = r + 1
    ws.cell(row=tot, column=1, value="全部合計")
    if lastr >= first:
        for col in "BCDFGHNO":
            ws[f"{col}{tot}"] = f"=SUM({col}{first}:{col}{lastr})"
        ws[f"E{tot}"] = f"=IF(B{tot}=0,0,C{tot}/B{tot})"
        ws[f"I{tot}"] = f"=IF(C{tot}=0,0,F{tot}/C{tot})"
        ws[f"J{tot}"] = f"=IF(D{tot}=0,0,G{tot}/D{tot})"
        ws[f"K{tot}"] = f"=MIN(K{first}:K{lastr})"
        ws[f"L{tot}"] = f'=IF(G{tot}=0,"",F{tot}/-G{tot})'
        ws[f"M{tot}"] = "n/a"
        ws[f"M{tot}"].comment = Comment("各作業回撤發生時間不同，不能直接相加", "report")

    for row in ws.iter_rows(min_row=first, max_row=tot, max_col=len(SUMMARY_HEADERS)):
        for c in row:
            c.font = Font(name=FONT, bold=(c.row == tot), color="0563C1" if c.column == 1 and c.row < tot else None,
                          underline="single" if c.column == 1 and c.row < tot else None)
            c.border = BORDER
            if c.row == tot:
                c.fill = TOT_FILL
            L = c.column_letter
            if L == "E":
                c.number_format = PCT
            elif L == "L":
                c.number_format = "0.00"
            elif L in "FGHIJKMNO":
                c.number_format = MONEY
    ws["G5"].comment = Comment("所有虧損單淨損益加總（負數）", "report")
    ws["M5"].comment = Comment("依平倉順序的累計損益，從高點回落的最大幅度", "report")
    ws["L5"].comment = Comment("總獲利 / |總虧損|；沒有虧損單時留空", "report")
    ws.freeze_panes = f"B{hr + 1}"
    set_widths(ws, [18, 9, 9, 9, 8, 12, 12, 12, 11, 11, 12, 9, 12, 11, 11])


OPEN_HEADERS = ["作業", "持倉ID", "商品", "方向", "手數", "開倉時間", "開倉價", "現價",
                "停損", "停利", "浮動損益", "庫存費", "註解", "Magic"]


def write_open(ws, open_pos, magic_map, patterns):
    ws.append(OPEN_HEADERS)
    style_header(ws, 1, len(OPEN_HEADERS))
    for p in sorted(open_pos, key=lambda x: x["time"]):
        ws.append([
            classify(p["magic"], p["comment"], magic_map, patterns), p["position_id"], p["symbol"],
            "Buy" if p["type"] == 0 else "Sell", p["volume"], ts(p["time"]), p["price_open"],
            p["price_current"], p["sl"], p["tp"], p["profit"], p["swap"], p["comment"], p["magic"],
        ])
    last = len(open_pos) + 1
    tot = last + 1
    ws.cell(row=tot, column=1, value="合計")
    ws[f"K{tot}"] = f"=SUM(K2:K{last})"
    ws[f"L{tot}"] = f"=SUM(L2:L{last})"
    for row in ws.iter_rows(min_row=2, max_row=tot, max_col=len(OPEN_HEADERS)):
        for c in row:
            c.font = Font(name=FONT, bold=(c.row == tot))
            c.border = BORDER
            if c.row == tot:
                c.fill = TOT_FILL
            if c.column_letter in "KL":
                c.number_format = MONEY
            elif c.column_letter == "F":
                c.number_format = TS
    set_widths(ws, [16, 12, 10, 6, 7, 19, 11, 11, 11, 11, 11, 10, 16, 12])


def build_workbook(rows, open_pos, account, start, end, currency, magic_map, patterns, out):
    grouped = OrderedDict()
    order = list(patterns) + sorted({r["strategy"] for r in rows} - set(patterns))
    for name in order:
        g = [r for r in rows if r["strategy"] == name]
        if g:
            grouped[name] = g

    wb = Workbook()
    summary = wb.active
    summary.title = "總覽"
    used = {"總覽", "未平倉", "全部明細"}
    sheets = OrderedDict()
    for name, g in grouped.items():
        sname = safe_sheet_name(name, used)
        last = write_strategy_sheet(wb.create_sheet(sname), name, g)
        sheets[name] = (sname, last)
    write_summary(summary, sheets, account, start, end, currency)

    allws = wb.create_sheet("全部明細")
    last = write_strategy_sheet(allws, "全部", rows)
    # 在最後一欄補上作業名稱（不用 insert_cols，以免公式欄位位移）
    col = len(STRAT_HEADERS) + 1
    allws.cell(row=1, column=col, value="作業")
    style_header(allws, 1, col)
    for i, r in enumerate(rows, 2):
        c = allws.cell(row=i, column=col, value=r["strategy"])
        c.font = Font(name=FONT)
        c.border = BORDER
    allws.column_dimensions[get_column_letter(col)].width = 14
    allws.auto_filter.ref = f"A1:{get_column_letter(col)}{last}"

    if open_pos:
        write_open(wb.create_sheet("未平倉"), open_pos, magic_map, patterns)
    wb.save(out)
    return grouped


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description="MT5 交易 → Excel（依作業分組計算損益）")
    ap.add_argument("--from", dest="start", default="2026-09-22", help="起始日期 YYYY-MM-DD（預設 2026-09-22）")
    ap.add_argument("--to", dest="end", default=None, help="結束日期 YYYY-MM-DD（預設：現在）")
    ap.add_argument("--strategies", nargs="*", default=["v5-3", "STT5M"],
                    help="作業名稱，開倉 comment 含有此字串就歸類（預設: v5-3 STT5M）")
    ap.add_argument("--map", nargs="*", default=[], metavar="MAGIC=NAME",
                    help="magic number 對應作業，例如 --map 1001=v5-3 2002=STT5M")
    ap.add_argument("--out", default=None, help="輸出檔名（預設 MT5_交易損益_<日期>.xlsx）")
    ap.add_argument("--csv", default=None, help="改從 CSV 讀成交資料（不連 MT5）")
    ap.add_argument("--dump", default=None, help="另存原始成交資料為 CSV")
    ap.add_argument("--login"), ap.add_argument("--password"), ap.add_argument("--server")
    ap.add_argument("--terminal-path", help="terminal64.exe 路徑（開多個 MT5 時指定）")
    a = ap.parse_args()

    start = datetime.strptime(a.start, "%Y-%m-%d")
    end = datetime.strptime(a.end, "%Y-%m-%d") + timedelta(days=1) if a.end else datetime.now() + timedelta(days=2)
    start_ts = int(start.replace(tzinfo=timezone.utc).timestamp())
    magic_map = {int(k): v for k, v in (m.split("=", 1) for m in a.map)}

    if a.csv:
        deals, open_pos, account, currency = load_from_csv(a.csv), [], "", ""
    else:
        # 往前多抓 30 天，好找回跨日持倉的進場成交
        deals, open_pos, account, currency = load_from_mt5(
            start - timedelta(days=30), end, a.login, a.password, a.server, a.terminal_path)
    if a.dump:
        dump_csv(deals, a.dump)
        print(f"原始成交已存: {a.dump}")

    end_ts = int(end.replace(tzinfo=timezone.utc).timestamp())
    rows = [r for r in build_positions(deals, start_ts, magic_map, a.strategies) if r["open_time"] < end_ts]
    out = a.out or f"MT5_交易損益_{start:%m%d}_{datetime.now():%m%d}.xlsx"
    grouped = build_workbook(rows, open_pos, account, start, min(end, datetime.now()), currency,
                             magic_map, a.strategies, out)

    print(f"\n已輸出: {out}")
    print(f"{'作業':<16}{'筆數':>6}{'虧損筆數':>9}{'總虧損':>14}{'淨損益':>14}")
    for name, g in grouped.items():
        net = [r["profit"] + r["commission"] + r["swap"] for r in g]
        loss = [x for x in net if x < 0]
        print(f"{name:<16}{len(g):>6}{len(loss):>9}{sum(loss):>14,.2f}{sum(net):>14,.2f}")
    if open_pos:
        print(f"未平倉 {len(open_pos)} 張，浮動損益 {sum(p['profit'] for p in open_pos):,.2f}")


if __name__ == "__main__":
    main()
